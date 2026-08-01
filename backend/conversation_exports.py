"""Exports portables et verifiables des conversations Marvin."""

from __future__ import annotations

import hashlib
import json
import os
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import database as db
from storage_layout import roots


FORMATS = {"md", "json", "docx", "pdf"}
MODES = {"lecture", "technique", "audit"}
_SECRET = re.compile(
    r"(?i)(authorization|api[_ -]?key|access[_ -]?token|refresh[_ -]?token|password|cookie)"
    r"\s*[:=]\s*([^\s,;]+)"
)


class ConversationExportError(ValueError):
    pass


@dataclass(frozen=True)
class ExportArtifact:
    path: Path
    format: str
    mode: str
    size_bytes: int
    sha256: str
    conversation_id: str

    def as_dict(self) -> dict[str, Any]:
        return {
            "success": True,
            "path": str(self.path),
            "name": self.path.name,
            "format": self.format,
            "mode": self.mode,
            "size_bytes": self.size_bytes,
            "sha256": self.sha256,
            "conversation_id": self.conversation_id,
        }


def _redact(value: Any) -> Any:
    if isinstance(value, str):
        return _SECRET.sub(lambda match: f"{match.group(1)}=[MASQUE]", value)
    if isinstance(value, list):
        return [_redact(item) for item in value]
    if isinstance(value, dict):
        return {
            key: "[MASQUE]" if re.search(r"(?i)password|token|secret|cookie|authorization", str(key)) else _redact(item)
            for key, item in value.items()
        }
    return value


def transcript(conversation_id: str, mode: str = "lecture") -> dict[str, Any]:
    if mode not in MODES:
        raise ConversationExportError(f"Mode invalide : {mode}")
    conversation = db.get_conversation(conversation_id)
    if not conversation:
        raise ConversationExportError("Conversation introuvable.")
    messages = []
    for message in db.get_messages_by_conversation(conversation_id, limit=1_000):
        if message.get("role") == "tool" and mode == "lecture":
            continue
        item = {
            "id": message.get("id"),
            "role": message.get("role"),
            "created_at": message.get("created_at"),
            "content": _redact(message.get("content") or ""),
        }
        if mode in {"technique", "audit"}:
            item["tool_calls"] = _redact(message.get("tool_calls") or [])
            item["tool_results"] = _redact(message.get("tool_results") or [])
        if mode == "audit":
            item["client_request_id"] = message.get("client_request_id")
        messages.append(item)
    payload = {
        "schema": "marvin.conversation-transcript.v1",
        "exported_at": datetime.now(timezone.utc).isoformat(),
        "mode": mode,
        "conversation": _redact(conversation),
        "messages": messages,
    }
    canonical = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    payload["integrity"] = {"algorithm": "sha256", "canonical_sha256": hashlib.sha256(canonical.encode("utf-8")).hexdigest()}
    return payload


def _safe_name(title: str) -> str:
    normalized = re.sub(r"[^\w.-]+", "-", title.strip(), flags=re.UNICODE).strip("-.")
    return (normalized or "conversation")[:80]


def _markdown(payload: dict[str, Any]) -> str:
    conversation = payload["conversation"]
    lines = [
        f"# {conversation.get('title') or 'Conversation Marvin'}",
        "",
        f"- Identifiant : `{conversation.get('id')}`",
        f"- Export : {payload['exported_at']}",
        f"- Mode : {payload['mode']}",
        f"- Empreinte : `{payload['integrity']['canonical_sha256']}`",
        "",
    ]
    labels = {"user": "Utilisateur", "assistant": "Marvin", "system": "Systeme", "tool": "Outil"}
    for message in payload["messages"]:
        lines.extend([f"## {labels.get(message['role'], message['role'])}", ""])
        if message.get("created_at"):
            lines.extend([f"_{message['created_at']}_", ""])
        lines.extend([str(message.get("content") or "_(aucun texte)_"), ""])
        if message.get("tool_calls"):
            lines.extend(["<details><summary>Appels d'outils</summary>", "", "```json", json.dumps(message["tool_calls"], ensure_ascii=False, indent=2), "```", "", "</details>", ""])
        if message.get("tool_results"):
            lines.extend(["<details><summary>Resultats d'outils</summary>", "", "```json", json.dumps(message["tool_results"], ensure_ascii=False, indent=2), "```", "", "</details>", ""])
    return "\n".join(lines).strip() + "\n"


def _atomic_bytes(path: Path, content: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_bytes(content)
    os.replace(temporary, path)


def _render_docx(path: Path, payload: dict[str, Any]) -> None:
    from docx import Document
    document = Document()
    document.add_heading(payload["conversation"].get("title") or "Conversation Marvin", 0)
    document.add_paragraph(f"Mode : {payload['mode']} | Export : {payload['exported_at']}")
    for message in payload["messages"]:
        document.add_heading({"user": "Utilisateur", "assistant": "Marvin", "tool": "Outil"}.get(message["role"], message["role"]), level=2)
        document.add_paragraph(str(message.get("content") or ""))
        if message.get("tool_calls"):
            document.add_paragraph(json.dumps(message["tool_calls"], ensure_ascii=False, indent=2), style="No Spacing")
        if message.get("tool_results"):
            document.add_paragraph(json.dumps(message["tool_results"], ensure_ascii=False, indent=2), style="No Spacing")
    temporary = path.with_suffix(".docx.tmp")
    document.save(str(temporary))
    os.replace(temporary, path)


def _render_pdf(path: Path, payload: dict[str, Any]) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape

    temporary = path.with_suffix(".pdf.tmp")
    doc = SimpleDocTemplate(str(temporary), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm)
    styles = getSampleStyleSheet()
    story = [Paragraph(escape(payload["conversation"].get("title") or "Conversation Marvin"), styles["Title"]), Spacer(1, 8)]
    for message in payload["messages"]:
        label = {"user": "Utilisateur", "assistant": "Marvin", "tool": "Outil"}.get(message["role"], message["role"])
        story.extend([Paragraph(escape(label), styles["Heading2"]), Paragraph(escape(str(message.get("content") or "")).replace("\n", "<br/>"), styles["BodyText"]), Spacer(1, 8)])
    doc.build(story)
    os.replace(temporary, path)


def export_conversation(conversation_id: str, format: str = "md", mode: str = "lecture") -> dict[str, Any]:
    selected_format = str(format).lower().lstrip(".")
    if selected_format not in FORMATS:
        raise ConversationExportError(f"Format invalide : {selected_format}")
    payload = transcript(conversation_id, mode)
    _, export_root = roots()
    directory = export_root / "documents" / "conversations"
    directory.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    path = directory / f"{conversation_id}-{_safe_name(str(payload['conversation'].get('title') or 'conversation'))}-{stamp}.{selected_format}"
    if selected_format == "json":
        _atomic_bytes(path, json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8"))
    elif selected_format == "md":
        _atomic_bytes(path, _markdown(payload).encode("utf-8"))
    elif selected_format == "docx":
        _render_docx(path, payload)
    else:
        _render_pdf(path, payload)
    content = path.read_bytes()
    if not content:
        raise ConversationExportError("L'artefact exporte est vide.")
    return ExportArtifact(path, selected_format, mode, len(content), hashlib.sha256(content).hexdigest(), conversation_id).as_dict()
