"""Manifeste de canevas éditable, stockage atomique et export PDF borné."""

from __future__ import annotations

import base64
import io
import json
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from storage_layout import roots

_ID = re.compile(r"^[a-zA-Z0-9_-]{1,80}$")
_HEX = re.compile(r"^#[0-9a-fA-F]{6}$")
_TYPES = {"text", "rect", "image", "line"}
MAX_LAYERS = 100
MAX_DATA_IMAGE = 2_000_000


def _directory() -> Path:
    path = roots()[1] / "documents" / "canvas"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _number(value: Any, low: float, high: float, default: float) -> float:
    try:
        return max(low, min(high, float(value)))
    except (TypeError, ValueError):
        return default


def _color(value: Any, default: str) -> str:
    candidate = str(value or "")
    return candidate.lower() if _HEX.fullmatch(candidate) else default


def validate_manifest(raw: dict[str, Any]) -> dict[str, Any]:
    if not isinstance(raw, dict):
        raise ValueError("Manifeste objet requis.")
    width = _number(raw.get("width"), 320, 2000, 794)
    height = _number(raw.get("height"), 320, 2800, 1123)
    layers = raw.get("layers")
    if not isinstance(layers, list) or len(layers) > MAX_LAYERS:
        raise ValueError(f"La liste layers est requise et limitée à {MAX_LAYERS} éléments.")
    normalized = []
    known = set()
    for index, item in enumerate(layers):
        if not isinstance(item, dict) or item.get("type") not in _TYPES:
            raise ValueError(f"Calque {index + 1} invalide.")
        layer_id = str(item.get("id") or f"layer-{index + 1}")
        if not _ID.fullmatch(layer_id) or layer_id in known:
            layer_id = f"layer-{index + 1}-{uuid.uuid4().hex[:6]}"
        known.add(layer_id)
        layer = {
            "id": layer_id, "type": item["type"], "name": str(item.get("name") or item["type"])[:80],
            "x": _number(item.get("x"), 0, width, 0), "y": _number(item.get("y"), 0, height, 0),
            "width": _number(item.get("width"), 1, width, 120), "height": _number(item.get("height"), 1, height, 60),
            "rotation": _number(item.get("rotation"), -360, 360, 0), "opacity": _number(item.get("opacity"), 0, 1, 1),
            "fill": _color(item.get("fill"), "#334155"), "color": _color(item.get("color"), "#172033"),
            "fontSize": _number(item.get("fontSize"), 6, 160, 24), "fontWeight": 700 if int(_number(item.get("fontWeight"), 100, 900, 400)) >= 600 else 400,
            "text": str(item.get("text") or "")[:5000], "locked": bool(item.get("locked", False)), "hidden": bool(item.get("hidden", False)),
        }
        source = str(item.get("src") or "")
        if source.startswith("data:image/") and len(source) <= MAX_DATA_IMAGE:
            layer["src"] = source
        normalized.append(layer)
    design_id = str(raw.get("id") or f"design-{uuid.uuid4().hex[:12]}")
    if not _ID.fullmatch(design_id):
        design_id = f"design-{uuid.uuid4().hex[:12]}"
    return {
        "schema_version": 1, "id": design_id, "title": str(raw.get("title") or "Document sans titre")[:160],
        "width": width, "height": height, "background": _color(raw.get("background"), "#ffffff"), "layers": normalized,
    }


def manifest_from_prompt(prompt: str, title: str = "") -> dict[str, Any]:
    """Turn a free visual brief into a safe editable starter manifest."""
    brief = str(prompt or "").strip()[:2000]
    if not brief:
        raise ValueError("Une description visuelle est requise.")
    from premium_document_design import resolve_design

    art = resolve_design({"design_prompt": brief})
    palette = art.palette
    heading = str(title or brief.splitlines()[0] or "Nouveau document")[:100]
    return validate_manifest({
        "title": heading, "background": palette["surface"], "width": 794, "height": 1123,
        "layers": [
            {"id": "accent", "type": "rect", "name": "Accent", "x": 0, "y": 0, "width": 26, "height": 1123, "fill": palette["accent"]},
            {"id": "title", "type": "text", "name": "Titre", "x": 78, "y": 92, "width": 620, "height": 120, "text": heading, "color": palette["text"], "fontSize": 54, "fontWeight": 700},
            {"id": "brief", "type": "text", "name": "Brief", "x": 82, "y": 245, "width": 590, "height": 150, "text": brief, "color": palette["muted"], "fontSize": 21},
            {"id": "rule", "type": "line", "name": "Séparateur", "x": 82, "y": 430, "width": 230, "height": 3, "fill": palette["secondary"]},
        ],
    })


def save(raw: dict[str, Any]) -> dict[str, Any]:
    manifest = validate_manifest(raw)
    target = _directory() / f"{manifest['id']}.json"
    temporary = target.with_suffix(".json.tmp")
    temporary.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(target)
    return {"success": True, "id": manifest["id"], "path": str(target), "manifest": manifest}


def load(design_id: str) -> dict[str, Any]:
    if not _ID.fullmatch(design_id):
        raise ValueError("Identifiant invalide.")
    target = _directory() / f"{design_id}.json"
    if not target.is_file():
        raise FileNotFoundError(design_id)
    return validate_manifest(json.loads(target.read_text(encoding="utf-8")))


def list_designs() -> list[dict[str, Any]]:
    result = []
    for path in sorted(_directory().glob("*.json"), key=lambda item: item.stat().st_mtime, reverse=True):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            result.append({"id": data.get("id"), "title": data.get("title"), "updated_at": path.stat().st_mtime, "layers": len(data.get("layers") or [])})
        except Exception:
            continue
    return result[:100]


def export_pdf(raw: dict[str, Any]) -> Path:
    manifest = validate_manifest(raw)
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader

    scale = 72 / 96
    target = _directory() / f"{manifest['id']}.pdf"
    pdf = canvas.Canvas(str(target), pagesize=(manifest["width"] * scale, manifest["height"] * scale))
    pdf.setFillColor(manifest["background"])
    pdf.rect(0, 0, manifest["width"] * scale, manifest["height"] * scale, fill=1, stroke=0)
    for layer in manifest["layers"]:
        if layer["hidden"]:
            continue
        x, y = layer["x"] * scale, (manifest["height"] - layer["y"] - layer["height"]) * scale
        w, h = layer["width"] * scale, layer["height"] * scale
        pdf.saveState(); pdf.setFillAlpha(layer["opacity"]); pdf.setStrokeAlpha(layer["opacity"])
        if layer["type"] == "rect":
            pdf.setFillColor(layer["fill"]); pdf.roundRect(x, y, w, h, 8 * scale, fill=1, stroke=0)
        elif layer["type"] == "line":
            pdf.setStrokeColor(layer["fill"]); pdf.setLineWidth(max(1, h)); pdf.line(x, y + h / 2, x + w, y + h / 2)
        elif layer["type"] == "text":
            pdf.setFillColor(layer["color"]); pdf.setFont("Helvetica-Bold" if layer["fontWeight"] >= 600 else "Helvetica", layer["fontSize"] * scale)
            text = pdf.beginText(x, y + h - layer["fontSize"] * scale)
            for line in layer["text"].splitlines() or [""]:
                text.textLine(line[:300])
            pdf.drawText(text)
        elif layer["type"] == "image" and layer.get("src", "").startswith("data:image/"):
            try:
                encoded = layer["src"].split(",", 1)[1]
                pdf.drawImage(ImageReader(io.BytesIO(base64.b64decode(encoded))), x, y, w, h, preserveAspectRatio=True, mask="auto")
            except Exception:
                pass
        pdf.restoreState()
    pdf.save()
    return target


def _image_bytes(layer: dict[str, Any]) -> io.BytesIO | None:
    source = layer.get("src", "")
    if not source.startswith("data:image/") or "," not in source:
        return None
    try:
        return io.BytesIO(base64.b64decode(source.split(",", 1)[1], validate=True))
    except (ValueError, base64.binascii.Error):
        return None


def export_pptx(raw: dict[str, Any]) -> Path:
    """Adapt the canvas to one editable PowerPoint slide."""
    manifest = validate_manifest(raw)
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width = Inches(manifest["width"] / 96)
    presentation.slide_height = Inches(manifest["height"] / 96)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.background.fill
    background.solid(); background.fore_color.rgb = RGBColor.from_string(manifest["background"][1:])
    for layer in manifest["layers"]:
        if layer["hidden"]:
            continue
        x, y, width, height = (Inches(layer[key] / 96) for key in ("x", "y", "width", "height"))
        if layer["type"] == "text":
            shape = slide.shapes.add_textbox(x, y, width, height)
            frame = shape.text_frame; frame.clear(); frame.word_wrap = True
            paragraph = frame.paragraphs[0]; paragraph.text = layer["text"]
            paragraph.font.size = Pt(layer["fontSize"] * .75)
            paragraph.font.bold = layer["fontWeight"] >= 600
            paragraph.font.color.rgb = RGBColor.from_string(layer["color"][1:])
        elif layer["type"] == "rect":
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, width, height)
            shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(layer["fill"][1:])
            shape.line.fill.background()
        elif layer["type"] == "line":
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, width, max(height, Pt(1)))
            shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(layer["fill"][1:])
            shape.line.fill.background()
        elif layer["type"] == "image":
            image = _image_bytes(layer)
            if image:
                slide.shapes.add_picture(image, x, y, width, height)
        if 'shape' in locals():
            shape.rotation = layer["rotation"]
            del shape
    target = _directory() / f"{manifest['id']}.pptx"
    presentation.save(target)
    return target


def export_docx(raw: dict[str, Any]) -> Path:
    """Adapt free-positioned content into an editable, ordered Word document."""
    manifest = validate_manifest(raw)
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(manifest["width"] / 96)
    section.page_height = Inches(manifest["height"] / 96)
    section.top_margin = section.bottom_margin = Inches(.55)
    section.left_margin = section.right_margin = Inches(.6)
    document.core_properties.title = manifest["title"]
    document.core_properties.comments = "Adaptation éditable générée depuis Marvin Studio"
    for layer in sorted((item for item in manifest["layers"] if not item["hidden"]), key=lambda item: (item["y"], item["x"])):
        if layer["type"] == "text" and layer["text"].strip():
            paragraph = document.add_paragraph()
            run = paragraph.add_run(layer["text"])
            run.bold = layer["fontWeight"] >= 600
            run.font.size = Pt(layer["fontSize"] * .75)
            run.font.color.rgb = RGBColor.from_string(layer["color"][1:])
        elif layer["type"] == "image":
            image = _image_bytes(layer)
            if image:
                paragraph = document.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(image, width=Inches(min(layer["width"] / 96, 6.8)))
        elif layer["type"] == "line":
            document.add_paragraph("─" * 48)
    document.add_paragraph().add_run(
        f"Adapté par Marvin Studio le {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
    ).italic = True
    target = _directory() / f"{manifest['id']}.docx"
    document.save(target)
    return target
