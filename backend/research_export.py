"""Exports bornés des rapports Research Engine, sans accès réseau."""
from __future__ import annotations

from io import BytesIO
import html
import json
import re
from typing import Any


_RUN_ID = re.compile(r"^swarm-[0-9a-f]{12}$")
_FORMATS = {"md", "json", "docx", "pdf"}


def _normalized_run(run: dict[str, Any]) -> tuple[str, str, list[dict[str, str]]]:
    run_id = str(run.get("id") or "")
    if not _RUN_ID.fullmatch(run_id):
        raise ValueError("Identifiant de recherche invalide.")
    answer = str(run.get("answer") or "")
    if not answer or len(answer) > 2_000_000:
        raise ValueError("Rapport absent ou trop volumineux.")
    sources = [
        {
            "id": str(source.get("id") or index + 1)[:80],
            "title": str(source.get("title") or source.get("url") or "")[:500],
            "url": str(source.get("url") or "")[:2_048],
        }
        for index, source in enumerate(list(run.get("sources") or [])[:200])
        if isinstance(source, dict)
    ]
    return run_id, answer, sources


def _markdown(answer: str, sources: list[dict[str, str]]) -> str:
    bibliography = "\n".join(
        f"- [{source['id']}] [{source['title']}]({source['url']})"
        for source in sources
    )
    return f"# Rapport de recherche Marvin\n\n{answer}\n\n## Sources\n\n{bibliography}\n"


def export_run(run: dict[str, Any], format_name: str) -> tuple[bytes, str, str]:
    requested = str(format_name or "").strip().lower()
    if requested not in _FORMATS:
        raise ValueError("Format invalide : md, json, docx ou pdf attendu.")
    run_id, answer, sources = _normalized_run(run)
    filename = f"marvin-research-{run_id}.{requested}"
    if requested == "json":
        payload = json.dumps(run, ensure_ascii=False, indent=2).encode("utf-8")
        return payload, "application/json; charset=utf-8", filename
    markdown = _markdown(answer, sources)
    if requested == "md":
        return markdown.encode("utf-8"), "text/markdown; charset=utf-8", filename
    if requested == "docx":
        from docx import Document

        document = Document()
        document.add_heading("Rapport de recherche Marvin", level=0)
        for line in answer.splitlines():
            stripped = line.strip()
            if not stripped:
                continue
            if stripped.startswith("### "):
                document.add_heading(stripped[4:], level=3)
            elif stripped.startswith("## "):
                document.add_heading(stripped[3:], level=2)
            elif stripped.startswith("# "):
                document.add_heading(stripped[2:], level=1)
            elif stripped.startswith(("- ", "* ")):
                document.add_paragraph(stripped[2:], style="List Bullet")
            else:
                document.add_paragraph(stripped)
        document.add_heading("Sources", level=1)
        for source in sources:
            document.add_paragraph(
                f"[{source['id']}] {source['title']} — {source['url']}",
                style="List Bullet",
            )
        output = BytesIO()
        document.save(output)
        return (
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename,
        )

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="Rapport de recherche Marvin",
    )
    styles = getSampleStyleSheet()
    story = [Paragraph("Rapport de recherche Marvin", styles["Title"]), Spacer(1, 8)]
    for line in answer.splitlines():
        stripped = line.strip()
        if not stripped:
            story.append(Spacer(1, 5))
            continue
        level = len(stripped) - len(stripped.lstrip("#"))
        clean = stripped[level:].strip() if level else stripped.lstrip("-* ").strip()
        style = styles["Heading2"] if level else styles["BodyText"]
        story.append(Paragraph(html.escape(clean), style))
        story.append(Spacer(1, 3))
    story.append(Paragraph("Sources", styles["Heading1"]))
    for source in sources:
        story.append(Paragraph(
            html.escape(f"[{source['id']}] {source['title']} — {source['url']}"),
            styles["BodyText"],
        ))
    document.build(story)
    return output.getvalue(), "application/pdf", filename
