import asyncio
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from tool_domains import documents
import libreoffice_runtime


class LetterDocxLayoutTests(unittest.TestCase):
    def test_letter_omits_placeholders_and_report_cover_metadata(self):
        with tempfile.TemporaryDirectory() as temp:
            target = Path(temp) / "lettre.docx"
            documents.configure(resolve_path=lambda value: str(Path(value).resolve()))
            with patch.object(libreoffice_runtime, "validate_docx_render", return_value={"success": True}):
                result = asyncio.run(documents.create_docx({
                    "path": str(target),
                    "title": "Lettre de motivation — Alternance Bachelor 3",
                    "subtitle": "Ynov Campus — [Ville] — [Date]",
                    "content": "Madame, Monsieur,\n\nJe vous présente ma candidature.",
                    "document_kind": "letter",
                    "author": "Chaouki BOUSSAYRI",
                    "letterhead": {"org": "KOIOS", "address": "[Adresse]", "email": "[Email]"},
                    "signature": {"lieu": "Aix-en-Provence", "name": "Chaouki BOUSSAYRI", "title": "Candidat Bachelor 3"},
                }))

            self.assertTrue(result["success"])
            word = Document(target)
            text = "\n".join(paragraph.text for paragraph in word.paragraphs)
            headers = "\n".join(paragraph.text for section in word.sections for paragraph in section.header.paragraphs)
            self.assertNotIn("[Ville]", text)
            self.assertNotIn("[Adresse]", headers)
            self.assertNotIn("[Email]", headers)
            self.assertNotIn("Chaouki BOUSSAYRI  ·", text)
            self.assertIn("Je vous présente ma candidature", text)


if __name__ == "__main__":
    unittest.main()
