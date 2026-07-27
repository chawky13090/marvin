"""Indexation RAG, questions documentaires et affinage ciblé de documents."""

from __future__ import annotations

import asyncio
import html
import json
import os
import re
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Awaitable, Callable, Protocol


class VectorStore(Protocol):
    async def index_chunks(
        self, doc_id: str, name: str, path: str, chunks: list[str]
    ) -> int: ...

    async def search(
        self, query: str, *, collections: list[str], limit: int
    ) -> list[dict]: ...

    def delete_doc_chunks(self, doc_id: str) -> None: ...


_resolver: Callable[[str], str] | None = None
_get_setting: Callable[[str], object] | None = None
_set_setting: Callable[[str, object], None] | None = None
_text_generator: Callable[..., Awaitable[str]] | None = None
_create_docx: Callable[[dict], Awaitable[dict]] | None = None
_vector_store: VectorStore | None = None
_generated_dir: Path | None = None

DOCUMENT_EXTENSIONS = {
    ".pdf", ".docx", ".txt", ".md", ".markdown", ".csv", ".tsv", ".log",
    ".json", ".xlsx", ".pptx", ".html", ".htm", ".rtf", ".xml", ".yaml",
    ".yml", ".ini", ".rst",
}
_MAX_DOCUMENTS = 200
_MAX_CHUNKS = 500
_MAX_RAG_K = 30


def configure(
    *,
    resolve_path: Callable[[str], str],
    get_setting: Callable[[str], object],
    set_setting: Callable[[str, object], None],
    text_generator: Callable[..., Awaitable[str]],
    create_docx: Callable[[dict], Awaitable[dict]],
    vector_store: VectorStore,
    generated_dir: Path,
) -> None:
    global _resolver, _get_setting, _set_setting, _text_generator
    global _create_docx, _vector_store, _generated_dir
    _resolver = resolve_path
    _get_setting = get_setting
    _set_setting = set_setting
    _text_generator = text_generator
    _create_docx = create_docx
    _vector_store = vector_store
    _generated_dir = generated_dir.resolve()


def _require_configured() -> None:
    if any(
        dependency is None
        for dependency in (
            _resolver,
            _get_setting,
            _set_setting,
            _text_generator,
            _create_docx,
            _vector_store,
            _generated_dir,
        )
    ):
        raise RuntimeError("Le domaine de connaissance documentaire n'est pas configuré")


def _resolve_path(raw: str) -> Path:
    _require_configured()
    return Path(_resolver(raw))  # type: ignore[misc]


def _indexed_docs() -> list[dict]:
    _require_configured()
    value = _get_setting("indexed_documents")  # type: ignore[misc]
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except json.JSONDecodeError:
            value = []
    return [dict(item) for item in value if isinstance(item, dict)] if isinstance(value, list) else []


def _save_indexed_docs(documents: list[dict]) -> None:
    _require_configured()
    _set_setting("indexed_documents", documents)  # type: ignore[misc]


def chunk_text(text: str, size: int = 900, overlap: int = 150) -> list[str]:
    if size <= 0 or overlap < 0 or overlap >= size:
        raise ValueError("Paramètres de découpage invalides")
    normalized = " ".join((text or "").split())
    if not normalized:
        return []
    step = size - overlap
    return [
        normalized[index:index + size]
        for index in range(0, len(normalized), step)
    ][:_MAX_CHUNKS]


def extract_text(path: Path) -> str:
    extension = path.suffix.lower()
    if extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise ValueError("PDF chiffré non indexable")
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if extension == ".docx":
        import docx

        document = docx.Document(str(path))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    if extension == ".pptx":
        from pptx import Presentation

        presentation = Presentation(str(path))
        return "\n".join(
            shape.text_frame.text
            for slide in presentation.slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
    if extension == ".xlsx":
        import openpyxl

        workbook = openpyxl.load_workbook(
            str(path), read_only=True, data_only=True
        )
        try:
            return "\n".join(
                " ".join(str(cell) for cell in row if cell is not None)
                for sheet in workbook.worksheets
                for row in sheet.iter_rows(values_only=True)
            )
        finally:
            workbook.close()
    raw = path.read_text(encoding="utf-8", errors="ignore")[:500_000]
    if extension in {".html", ".htm"}:
        raw = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", raw)
        raw = re.sub(r"(?s)<[^>]+>", " ", raw)
        raw = html.unescape(raw)
    elif extension == ".rtf":
        raw = re.sub(r"\\'[0-9a-fA-F]{2}", "", raw)
        raw = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", raw)
        raw = raw.replace("{", " ").replace("}", " ")
    return raw


def collect_targets(directory: Path) -> list[Path]:
    targets: list[Path] = []
    for root, directories, files in os.walk(directory):
        directories.sort()
        for filename in sorted(files):
            candidate = Path(root) / filename
            if candidate.suffix.lower() in DOCUMENT_EXTENSIONS:
                targets.append(candidate)
                if len(targets) >= _MAX_DOCUMENTS:
                    return targets
    return targets


async def index(args: dict) -> dict:
    raw_path = str(args.get("path") or "").strip()
    if not raw_path:
        return {"error": "Indiquez le chemin du fichier ou dossier à indexer."}
    try:
        path = _resolve_path(raw_path)
        if not path.exists():
            return {"error": f"Introuvable : {path}"}
        if path.is_dir():
            targets = await asyncio.to_thread(collect_targets, path)
        elif path.suffix.lower() in DOCUMENT_EXTENSIONS:
            targets = [path]
        else:
            return {
                "error": (
                    "Type non pris en charge. Formats : "
                    f"{', '.join(sorted(DOCUMENT_EXTENSIONS))}"
                )
            }
        if not targets:
            return {"error": "Aucun document compatible trouvé."}

        documents = _indexed_docs()
        by_path = {str(item.get("path")): item for item in documents}
        try:
            concurrency = int(_get_setting("task_concurrency") or 3)  # type: ignore[misc]
        except (TypeError, ValueError):
            concurrency = 3
        concurrency = min(max(concurrency, 1), 16)
        semaphore = asyncio.Semaphore(concurrency)

        async def extract(target: Path):
            async with semaphore:
                try:
                    text = await asyncio.to_thread(extract_text, target)
                    return target, chunk_text(text)
                except (OSError, RuntimeError, TypeError, ValueError) as exc:
                    return target, exc

        extracted = await asyncio.gather(*(extract(target) for target in targets))
        results: list[str] = []
        failures: list[dict] = []
        total_chunks = 0
        for target, chunks in extracted:
            if isinstance(chunks, Exception):
                failures.append({"name": target.name, "error": str(chunks)[:160]})
                continue
            if not chunks:
                failures.append({"name": target.name, "error": "document vide"})
                continue
            try:
                previous = by_path.get(str(target), {})
                document_id = str(previous.get("id") or uuid.uuid4().hex[:12])
                if previous:
                    _vector_store.delete_doc_chunks(document_id)  # type: ignore[union-attr]
                count = await _vector_store.index_chunks(  # type: ignore[union-attr]
                    document_id, target.name, str(target), chunks
                )
                total_chunks += count
                by_path[str(target)] = {
                    "id": document_id,
                    "name": target.name,
                    "path": str(target),
                    "chunks": count,
                    "indexed_at": datetime.now(timezone.utc).isoformat(),
                }
                results.append(target.name)
            except (OSError, RuntimeError, TypeError, ValueError) as exc:
                failures.append({"name": target.name, "error": str(exc)[:160]})
        _save_indexed_docs(list(by_path.values()))
        return {
            "success": True,
            "indexed": results,
            "documents": len(results),
            "failed": failures,
            "failed_count": len(failures),
            "chunks": total_chunks,
            "message": (
                f"{len(results)} document(s) indexé(s) ({total_chunks} fragments). "
                "Tu peux maintenant poser des questions dessus."
            ),
        }
    except (OSError, RuntimeError, TypeError, ValueError) as exc:
        return {"error": str(exc)}


async def ask(args: dict) -> dict:
    question = str(args.get("question") or args.get("query") or "").strip()
    if not question:
        return {"error": "Pose une question (paramètre 'question')."}
    try:
        if not _indexed_docs():
            return {"error": "Aucun document indexé. Indexe d'abord un fichier ou un dossier."}
        try:
            limit = int(args.get("k") or _get_setting("rag_k") or 6)  # type: ignore[misc]
        except (TypeError, ValueError):
            limit = 6
        limit = min(max(limit, 1), _MAX_RAG_K)
        seen_texts: set[str] = set()
        all_hits: list[dict] = []

        async def append_hits(query: str) -> None:
            hits = await _vector_store.search(  # type: ignore[union-attr]
                query, collections=["documents"], limit=limit
            )
            for hit in hits:
                text = str(hit.get("text") or "")
                key = text[:200]
                if key and key not in seen_texts:
                    seen_texts.add(key)
                    all_hits.append(hit)

        current_query = question
        await append_hits(question)
        if all_hits:
            digest = "\n\n".join(
                f"[{index + 1}] ({hit.get('metadata', {}).get('name', '?')})\n"
                f"{str(hit.get('text') or '')[:600]}"
                for index, hit in enumerate(all_hits[:8])
            )
            judge_prompt = (
                f"Question : {question}\n\nExtraits déjà trouvés :\n{digest}\n\n"
                "Réponds uniquement SUFFISANT, ou par une requête courte complémentaire."
            )
            try:
                verdict = str(
                    await _text_generator(  # type: ignore[misc]
                        judge_prompt, model=args.get("model"), timeout=45
                    )
                    or ""
                ).strip()
            except (OSError, RuntimeError, TypeError, ValueError):
                verdict = "SUFFISANT"
            if (
                verdict
                and not verdict.upper().startswith("SUFFISANT")
                and len(verdict) <= 150
                and verdict.lower() != question.lower()
            ):
                await append_hits(verdict)
                current_query = verdict
        if not all_hits:
            return {
                "success": True,
                "answer": (
                    "Je n'ai trouvé aucun passage pertinent dans tes documents "
                    "pour cette question."
                ),
                "sources": [],
                "message": "Aucun passage pertinent.",
            }

        context = "\n\n".join(
            f"[{index + 1}] (source : {hit.get('metadata', {}).get('name', '?')})\n"
            f"{hit.get('text', '')}"
            for index, hit in enumerate(all_hits[: limit * 2])
        )
        prompt = (
            "Réponds en français uniquement à partir des extraits et cite le fichier. "
            "Si l'information manque, dis-le.\n\n"
            f"EXTRAITS :\n{context}\n\nQUESTION : {question}\n\nRÉPONSE :"
        )
        answer = await _text_generator(  # type: ignore[misc]
            prompt, model=args.get("model"), timeout=180
        )
        sources: list[dict] = []
        seen_names: set[str] = set()
        for hit in all_hits:
            name = str(hit.get("metadata", {}).get("name", "?"))
            if name not in seen_names:
                seen_names.add(name)
                sources.append(
                    {
                        "name": name,
                        "score": hit.get("score"),
                        "extrait": str(hit.get("text") or "")[:200],
                    }
                )
        return {
            "success": True,
            "answer": str(answer or "").strip(),
            "sources": sources,
            "message": (
                f"Réponse basée sur {len(sources)} document(s)"
                + (
                    " (2 passes de recherche)."
                    if current_query != question
                    else "."
                )
            ),
        }
    except (OSError, RuntimeError, TypeError, ValueError) as exc:
        return {"error": str(exc)}


async def list_documents(args: dict) -> dict:
    try:
        documents = _indexed_docs()
        return {"success": True, "documents": documents, "count": len(documents)}
    except (RuntimeError, TypeError, ValueError) as exc:
        return {"error": str(exc)}


async def forget(args: dict) -> dict:
    key = str(args.get("id") or args.get("name") or "").strip().lower()
    if not key:
        return {"error": "Indique l'id ou le nom du document à oublier."}
    try:
        documents = _indexed_docs()
        kept: list[dict] = []
        removed: dict | None = None
        for document in documents:
            if (
                str(document.get("id") or "").lower() == key
                or str(document.get("name") or "").lower() == key
            ):
                removed = document
            else:
                kept.append(document)
        if removed is None:
            return {"error": "Document introuvable dans l'index."}
        _vector_store.delete_doc_chunks(str(removed["id"]))  # type: ignore[union-attr]
        _save_indexed_docs(kept)
        return {
            "success": True,
            "message": f"« {removed.get('name')} » retiré de l'index.",
        }
    except (OSError, RuntimeError, TypeError, ValueError) as exc:
        return {"error": str(exc)}


async def refine_section(args: dict) -> dict:
    raw_file = str(args.get("file") or "").strip()
    section = str(args.get("section") or "").strip()
    instruction = str(
        args.get("instruction") or "développe et complète cette section"
    ).strip()
    notes = str(args.get("notes") or "").strip()
    if not raw_file or not section:
        return {
            "error": (
                "Fournis 'file' (le .docx généré) et 'section' "
                "(le titre à affiner)."
            )
        }
    try:
        _require_configured()
        base = _generated_dir.resolve()  # type: ignore[union-attr]
        requested = _resolve_path(raw_file).resolve()
        try:
            requested.relative_to(base)
            document_path = requested
        except ValueError:
            document_path = (base / Path(raw_file).name).resolve()
            document_path.relative_to(base)
        metadata_path = Path(f"{document_path}.meta.json")
        if not metadata_path.is_file():
            return {
                "error": (
                    "Document non affinable (régénère-le d'abord avec "
                    "la nouvelle version)."
                )
            }
        try:
            metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            return {"error": f"Métadonnées illisibles : {str(exc)[:80]}"}
        markdown = str(metadata.get("markdown") or "")
        pattern = re.compile(
            r"(^|\n)(#{2,3})\s*" + re.escape(section) + r"\s*\n",
            re.IGNORECASE,
        )
        match = pattern.search(markdown)
        if not match:
            candidates = metadata.get("sections", [])
            candidate = next(
                (
                    value
                    for value in candidates
                    if section.lower() in str(value).lower()
                    or str(value).lower() in section.lower()
                ),
                None,
            )
            if candidate:
                section = str(candidate)
                pattern = re.compile(
                    r"(^|\n)(#{2,3})\s*" + re.escape(section) + r"\s*\n",
                    re.IGNORECASE,
                )
                match = pattern.search(markdown)
        if not match:
            return {
                "error": (
                    f"Section « {section} » introuvable. Sections : "
                    f"{', '.join(map(str, metadata.get('sections', [])))}"
                )
            }
        start = match.start() if match.start() == 0 else match.start() + 1
        following = re.search(r"\n#{1,3}\s", markdown[match.end():])
        end = match.end() + following.start() if following else len(markdown)
        current = markdown[start:end].strip()
        prompt = (
            f"Réécris uniquement la section « {section} » en français professionnel.\n"
            f"SECTION ACTUELLE :\n{current}\n\nCONSIGNE : {instruction}\n"
            + (f"ÉLÉMENTS À INTÉGRER :\n{notes}\n" if notes else "")
            + f"Commence par « ## {section} »."
        )
        new_section = str(
            await _text_generator(  # type: ignore[misc]
                prompt, model=args.get("model"), timeout=240
            )
            or ""
        ).strip()
        if not new_section:
            return {"error": "La régénération n'a rien produit."}
        if "##" not in new_section[:80]:
            new_section = f"## {section}\n\n{new_section}"
        new_markdown = re.sub(
            r"\n{3,}",
            "\n\n",
            markdown[:start] + new_section + "\n\n" + markdown[end:],
        )
        rendered = await _create_docx(  # type: ignore[misc]
            {
                "path": str(document_path),
                "title": metadata.get("title", "Document"),
                "subtitle": metadata.get("subtitle", ""),
                "content": new_markdown,
                "author": metadata.get("author", "Marvin"),
                "letterhead": metadata.get("letterhead") or {},
                "references": metadata.get("references") or {},
                "signature": metadata.get("signature") or {},
            }
        )
        if rendered.get("error"):
            return {"error": f"Rendu Word échoué : {rendered['error']}"}
        metadata["markdown"] = new_markdown
        temporary = metadata_path.with_name(
            f".{metadata_path.name}.{os.getpid()}.{time.time_ns()}.tmp"
        )
        try:
            temporary.write_text(
                json.dumps(metadata, ensure_ascii=False), encoding="utf-8"
            )
            os.replace(temporary, metadata_path)
        finally:
            temporary.unlink(missing_ok=True)
        return {
            "success": True,
            "file": str(document_path),
            "section": section,
            "section_markdown": new_section[:3000],
            "markdown": new_markdown[:6000],
            "sections": metadata.get("sections", []),
            "message": f"Section « {section} » affinée → {document_path.name}",
        }
    except (OSError, RuntimeError, TypeError, ValueError) as exc:
        return {"error": str(exc)}
