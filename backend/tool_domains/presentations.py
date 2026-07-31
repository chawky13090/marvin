"""Création et lecture bornées des présentations PowerPoint."""

from __future__ import annotations

import asyncio
import hashlib
import json
import os
import tempfile
import time
from pathlib import Path
from typing import Any, Callable

from office_packages import validate_office_package
from platform_paths import PATHS as _SYSTEM_PATHS


PathResolver = Callable[[str], str]

_MAX_SLIDES = 100
_MAX_BULLETS = 100
_MAX_FIELD_CHARS = 10_000
_MAX_READ_SLIDES = 500
_MAX_READ_TEXT = 100_000
_resolver: PathResolver | None = None


def configure(*, resolve_path: PathResolver) -> None:
    global _resolver
    _resolver = resolve_path


def _resolve_path(value: str) -> str:
    if _resolver is None:
        raise RuntimeError("Le domaine présentations n'est pas configuré.")
    return _resolver(value)


def _text(value: object, *, name: str, maximum: int = _MAX_FIELD_CHARS) -> str:
    text = "" if value is None else str(value)
    if len(text) > maximum:
        raise ValueError(f"Champ '{name}' limité à {maximum} caractères.")
    return text


def _slides(value: object) -> list[dict[str, Any]]:
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except json.JSONDecodeError as exc:
            raise ValueError(
                "'slides' doit être une liste d'objets {title, bullets}."
            ) from exc
    if value in (None, ""):
        return []
    if not isinstance(value, list):
        raise ValueError("'slides' doit être une liste.")
    if len(value) > _MAX_SLIDES:
        raise ValueError(f"Présentation limitée à {_MAX_SLIDES} diapositives de contenu.")
    normalized: list[dict[str, Any]] = []
    for index, slide in enumerate(value):
        if isinstance(slide, str):
            slide = {"title": slide, "bullets": []}
        if not isinstance(slide, dict):
            raise ValueError(f"Diapositive {index + 1} invalide.")
        title = _text(slide.get("title"), name=f"slides[{index}].title")
        bullets = slide.get("bullets")
        if bullets is None:
            bullets = slide.get("points", [])
        if isinstance(bullets, str):
            bullets = [item.strip() for item in bullets.splitlines() if item.strip()]
        if not isinstance(bullets, list) or any(
            not isinstance(item, (str, int, float)) or isinstance(item, bool)
            for item in bullets
        ):
            raise ValueError(f"Puces de la diapositive {index + 1} invalides.")
        if len(bullets) > _MAX_BULLETS:
            raise ValueError(
                f"Diapositive {index + 1} limitée à {_MAX_BULLETS} puces."
            )
        normalized.append(
            {
                "title": title,
                "bullets": [
                    _text(item, name=f"slides[{index}].bullets") for item in bullets
                ],
            }
        )
    return normalized


def _create_presentation(
    *,
    title: str,
    subtitle: str,
    slides: list[dict[str, Any]],
    output: Path,
    style: str = "professional_dark",
) -> int:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    presentation = Presentation()
    if style == "professional_dark":
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        background = RGBColor(0x12, 0x14, 0x18)
        panel = RGBColor(0x1B, 0x1E, 0x24)
        foreground = RGBColor(0xF0, 0xF0, 0xEC)
        muted = RGBColor(0xA6, 0xA7, 0xA9)
        accent = RGBColor(0x91, 0xA3, 0x98)

        def add_text(slide: Any, left: float, top: float, width: float,
                     height: float, text: str, *, size: int, color: Any,
                     bold: bool = False, align: Any = PP_ALIGN.LEFT) -> None:
            box = slide.shapes.add_textbox(
                Inches(left), Inches(top), Inches(width), Inches(height)
            )
            frame = box.text_frame
            frame.word_wrap = True
            paragraph = frame.paragraphs[0]
            paragraph.text = text
            paragraph.alignment = align
            paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = color

        def prepare(slide: Any, *, cover: bool = False) -> None:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = background if cover else panel
            bar = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, 0, 0,
                Inches(0.10 if cover else 13.333),
                Inches(7.5 if cover else 0.06),
            )
            bar.fill.solid()
            bar.fill.fore_color.rgb = accent
            bar.line.fill.background()

        if title:
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            prepare(slide, cover=True)
            add_text(slide, 1.25, 1.45, 10.7, 1.35, title, size=48,
                     color=foreground, bold=True)
            line = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(1.27), Inches(2.90),
                Inches(2.35), Inches(0.04)
            )
            line.fill.solid()
            line.fill.fore_color.rgb = accent
            line.line.fill.background()
            if subtitle:
                add_text(slide, 1.25, 3.25, 10.5, 1.1, subtitle, size=21,
                         color=muted)
        for number, definition in enumerate(slides, start=2 if title else 1):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            prepare(slide)
            add_text(slide, 1.05, 0.72, 10.9, 0.75, definition["title"],
                     size=30, color=foreground, bold=True)
            bullets = [item for item in definition["bullets"] if item.strip()][:8]
            top = 1.75
            spacing = min(0.64, 4.75 / max(len(bullets), 1))
            for bullet in bullets:
                add_text(slide, 1.18, top, 0.28, 0.42, "•", size=18,
                         color=accent, bold=True)
                add_text(slide, 1.58, top, 10.35, max(0.42, spacing), bullet,
                         size=17, color=foreground)
                top += spacing
            add_text(slide, 11.65, 7.05, 0.55, 0.25, str(number), size=9,
                     color=muted, align=PP_ALIGN.RIGHT)
    else:
        if title:
            slide = presentation.slides.add_slide(presentation.slide_layouts[0])
            if slide.shapes.title is not None:
                slide.shapes.title.text = title
            if subtitle and len(slide.placeholders) > 1:
                slide.placeholders[1].text = subtitle
        for definition in slides:
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            if slide.shapes.title is not None:
                slide.shapes.title.text = definition["title"]
            if len(slide.placeholders) > 1:
                frame = slide.placeholders[1].text_frame
                bullets = definition["bullets"]
                if bullets:
                    frame.text = bullets[0]
                    for bullet in bullets[1:]:
                        paragraph = frame.add_paragraph()
                        paragraph.text = bullet
                else:
                    frame.clear()
    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output.stem}.",
        suffix=".pptx",
        dir=str(output.parent),
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        presentation.save(temporary)
        validate_office_package(
            temporary,
            suffix=".pptx",
            marker="ppt/presentation.xml",
        )
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)
    return len(presentation.slides)


async def create(args: dict[str, Any]) -> dict[str, Any]:
    try:
        title = _text(args.get("title"), name="title")
        subtitle = _text(args.get("subtitle"), name="subtitle")
        slides = _slides(args.get("slides"))
    except ValueError as exc:
        return {"error": str(exc)}
    if not title and not slides:
        return {"error": "Fournir 'title' et/ou 'slides' (liste de {title, bullets})."}
    output_value = str(args.get("output_path") or "").strip()
    style = str(args.get("style") or "professional_dark").strip().casefold()
    if style not in {"professional_dark", "office"}:
        return {"error": "Style invalide : utiliser 'professional_dark' ou 'office'."}
    try:
        output = (
            Path(_resolve_path(output_value))
            if output_value
            else Path(_SYSTEM_PATHS["desktop"]) / "presentation.pptx"
        )
    except (OSError, RuntimeError, ValueError) as exc:
        return {"error": f"Chemin de sortie invalide : {exc}"}
    if output_value.casefold() in {"bureau", "desktop"} or output.is_dir():
        output = output / "Presentation_Marvin.pptx"
    if output.suffix.casefold() != ".pptx":
        return {"error": "Le fichier de sortie doit utiliser l'extension .pptx."}
    try:
        count = await asyncio.to_thread(
            _create_presentation,
            title=title,
            subtitle=subtitle,
            slides=slides,
            output=output,
            style=style,
        )
    except PermissionError:
        output = output.with_name(
            f"{output.stem}-nouvelle-version-{int(time.time())}{output.suffix}"
        )
        try:
            count = await asyncio.to_thread(
                _create_presentation,
                title=title,
                subtitle=subtitle,
                slides=slides,
                output=output,
                style=style,
            )
        except (OSError, ValueError) as exc:
            return {"error": f"Création PowerPoint échouée : {exc}"}
    except (OSError, ValueError) as exc:
        return {"error": f"Création PowerPoint échouée : {exc}"}
    return {
        "success": True,
        "verified": True,
        "path": str(output),
        "destination": str(output),
        "slides": count,
        "style": style,
        "size_bytes": output.stat().st_size,
        "sha256": await asyncio.to_thread(
            lambda: hashlib.sha256(output.read_bytes()).hexdigest()
        ),
        "artifacts": [
            {
                "path": str(output),
                "type": "presentation",
                "format": "pptx",
                "slides": count,
            }
        ],
        "message": f"Présentation créée ({count} diapos) → {output}",
    }


def _read_presentation(path: Path) -> tuple[list[dict[str, Any]], str, bool]:
    from pptx import Presentation

    validate_office_package(path, suffix=".pptx", marker="ppt/presentation.xml")
    presentation = Presentation(path)
    content: list[dict[str, Any]] = []
    full_parts: list[str] = []
    total_chars = 0
    truncated = False
    for index, slide in enumerate(presentation.slides):
        if index >= _MAX_READ_SLIDES:
            truncated = True
            break
        parts: list[str] = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            remaining = _MAX_READ_TEXT - total_chars
            if remaining <= 0:
                truncated = True
                break
            if len(text) > remaining:
                text = text[:remaining]
                truncated = True
            parts.append(text)
            total_chars += len(text)
        slide_text = "\n".join(parts)
        content.append({"slide": index + 1, "text": slide_text})
        full_parts.append(f"--- Diapo {index + 1} ---\n{slide_text}")
        if truncated:
            break
    return content, "\n\n".join(full_parts), truncated


async def read(args: dict[str, Any]) -> dict[str, Any]:
    value = args.get("path")
    if not isinstance(value, str) or not value.strip():
        return {"error": "Paramètre 'path' requis."}
    try:
        path = Path(_resolve_path(value.strip()))
        content, full, truncated = await asyncio.to_thread(_read_presentation, path)
    except FileNotFoundError as exc:
        return {"error": str(exc)}
    except (OSError, ValueError) as exc:
        return {"error": f"Lecture PowerPoint échouée : {exc}"}
    return {
        "success": True,
        "slides": len(content),
        "content": content,
        "text": full[:6_000],
        "truncated": truncated or len(full) > 6_000,
        "message": f"{len(content)} diapo(s) lue(s)",
    }
