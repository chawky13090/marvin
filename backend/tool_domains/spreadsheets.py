"""Création, lecture et conversion des tableurs de Marvin."""

from __future__ import annotations

import csv
import os
import time
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from platform_paths import PATHS as _SYSTEM_PATHS


_resolver: Callable[[str], str] | None = None
_MAX_READ_ROWS = 10_000
_MAX_CSV_ROWS = 1_000_000
_INVALID_SHEET_CHARS = frozenset(r"[]:*?/\\")


def configure(*, resolve_path: Callable[[str], str]) -> None:
    global _resolver
    _resolver = resolve_path


def _resolve_path(raw: str) -> str:
    if _resolver is None:
        raise RuntimeError("Le domaine tableurs n'est pas configuré")
    return _resolver(raw)


def _sheet_name(raw: object, fallback: str = "Feuille1") -> str:
    name = str(raw or fallback).strip() or fallback
    if any(char in _INVALID_SHEET_CHARS for char in name):
        raise ValueError("Le nom de feuille contient un caractère interdit : []:*?/\\")
    return name[:31]


def _delimiter(raw: object, default: str = ",") -> str:
    value = str(raw or default)
    if len(value) != 1:
        raise ValueError("Le délimiteur CSV doit contenir exactement un caractère.")
    return value


def _atomic_save(workbook: Workbook, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(
        f".{destination.stem}.{os.getpid()}.{time.time_ns()}{destination.suffix}"
    )
    try:
        workbook.save(temporary)
        os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)


def _normalize_rows(raw: object) -> Iterable[list[object]]:
    if raw is None:
        return []
    if not isinstance(raw, list):
        raise ValueError("Le paramètre 'rows' doit être une liste de lignes.")
    for row in raw:
        if not isinstance(row, (list, tuple)):
            raise ValueError("Chaque ligne doit être une liste de cellules.")
        yield list(row)


async def create_xlsx(args: dict) -> dict:
    """Crée un fichier Excel (.xlsx) avec des données."""
    workbook: Workbook | None = None
    try:
        raw_path = str(args.get("path") or "").strip()
        path = (
            _resolve_path(raw_path)
            if raw_path
            else str(
                Path(_SYSTEM_PATHS["desktop"]) / f"tableau_{int(time.time())}.xlsx"
            )
        )
        title = str(args.get("title") or "")
        headers_raw = args.get("headers", [])
        if not isinstance(headers_raw, list):
            return {"error": "Le paramètre 'headers' doit être une liste."}
        headers = list(headers_raw)
        rows = list(_normalize_rows(args.get("rows", [])))
        sheet_name = _sheet_name(args.get("sheet_name"))
        column_count = max(
            [len(headers), *(len(row) for row in rows)] if rows else [len(headers), 1]
        )

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name

        if title:
            sheet.merge_cells(
                start_row=1, start_column=1, end_row=1, end_column=column_count
            )
            cell = sheet.cell(row=1, column=1, value=title)
            cell.font = Font(bold=True, color="FFFFFF", size=13)
            cell.fill = PatternFill("solid", fgColor="C9A84C")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            sheet.row_dimensions[1].height = 26
            start_row = 2
        else:
            start_row = 1

        for column, header in enumerate(headers, 1):
            cell = sheet.cell(row=start_row, column=column, value=header)
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.fill = PatternFill("solid", fgColor="2A2A2A")
            cell.alignment = Alignment(horizontal="center")
        if headers:
            start_row += 1

        for row_index, row in enumerate(rows):
            fill_color = "181818" if row_index % 2 == 0 else "111111"
            for column, value in enumerate(row, 1):
                cell = sheet.cell(row=start_row + row_index, column=column, value=value)
                cell.fill = PatternFill("solid", fgColor=fill_color)
                cell.font = Font(color="C3C3C3", size=10)
                cell.alignment = Alignment(horizontal="left")

        for column in range(1, sheet.max_column + 1):
            max_length = max(
                (
                    len(str(sheet.cell(row=row, column=column).value or ""))
                    for row in range(1, sheet.max_row + 1)
                ),
                default=8,
            )
            sheet.column_dimensions[get_column_letter(column)].width = min(
                max_length + 4, 40
            )

        output = Path(path)
        if not output.suffix:
            output = output.with_suffix(".xlsx")
        if output.suffix.lower() != ".xlsx":
            return {"error": "Le fichier de sortie doit utiliser l'extension .xlsx."}
        _atomic_save(workbook, output)
        return {
            "success": True,
            "message": f"Fichier Excel créé : {output}",
            "path": str(output),
        }
    except (OSError, TypeError, ValueError) as exc:
        return {"error": str(exc)}
    finally:
        if workbook is not None:
            workbook.close()


async def csv_to_xlsx(args: dict) -> dict:
    """Convertit un fichier CSV en classeur Excel avec traitement borné en flux."""
    workbook: Workbook | None = None
    try:
        raw_path = str(args.get("path") or "").strip()
        if not raw_path:
            return {"error": "Paramètre 'path' requis."}
        source = Path(_resolve_path(raw_path))
        if not source.is_file():
            return {"error": f"Fichier introuvable : {source}"}

        raw_delimiter = args.get("delimiter")
        with source.open("r", encoding="utf-8-sig", newline="") as handle:
            sample = handle.read(2048)
            handle.seek(0)
            if raw_delimiter:
                delimiter = _delimiter(raw_delimiter)
            else:
                try:
                    delimiter = csv.Sniffer().sniff(
                        sample, delimiters=",;\t|"
                    ).delimiter
                except csv.Error:
                    delimiter = ","

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = _sheet_name(source.stem)
            row_count = 0
            for row_count, row in enumerate(
                csv.reader(handle, delimiter=delimiter), start=1
            ):
                if row_count > _MAX_CSV_ROWS:
                    raise ValueError(
                        f"CSV refusé : plus de {_MAX_CSV_ROWS} lignes"
                    )
                for column, value in enumerate(row, 1):
                    cell = sheet.cell(row=row_count, column=column, value=value)
                    if row_count == 1:
                        cell.font = Font(bold=True)

        output_raw = str(args.get("output_path") or "").strip()
        output = Path(
            _resolve_path(output_raw or str(source.with_suffix(".xlsx")))
        )
        if output.suffix.lower() != ".xlsx":
            return {"error": "Le fichier de sortie doit utiliser l'extension .xlsx."}
        _atomic_save(workbook, output)
        return {
            "success": True,
            "path": str(output),
            "destination": str(output),
            "rows": row_count,
            "message": f"CSV converti en Excel ({row_count} lignes) → {output}",
        }
    except (OSError, TypeError, ValueError, csv.Error) as exc:
        return {"error": str(exc)}
    finally:
        if workbook is not None:
            workbook.close()


async def xlsx_to_csv(args: dict) -> dict:
    """Convertit une feuille d'un classeur Excel (.xlsx) en CSV."""
    workbook = None
    try:
        raw_path = str(args.get("path") or "").strip()
        if not raw_path:
            return {"error": "Paramètre 'path' requis."}
        source = Path(_resolve_path(raw_path))
        if not source.is_file():
            return {"error": f"Fichier introuvable : {source}"}
        delimiter = _delimiter(args.get("delimiter"), ",")

        workbook = load_workbook(str(source), read_only=True, data_only=True)
        requested_sheet = str(args.get("sheet") or "")
        sheet = (
            workbook[requested_sheet]
            if requested_sheet and requested_sheet in workbook.sheetnames
            else workbook.active
        )
        output_raw = str(args.get("output_path") or "").strip()
        output = Path(_resolve_path(output_raw or str(source.with_suffix(".csv"))))
        output.parent.mkdir(parents=True, exist_ok=True)
        temporary = output.with_name(
            f".{output.name}.{os.getpid()}.{time.time_ns()}.tmp"
        )
        row_count = 0
        try:
            with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle, delimiter=delimiter)
                for row in sheet.iter_rows(values_only=True):
                    writer.writerow(["" if value is None else value for value in row])
                    row_count += 1
            os.replace(temporary, output)
        finally:
            temporary.unlink(missing_ok=True)
        return {
            "success": True,
            "path": str(output),
            "destination": str(output),
            "rows": row_count,
            "sheet": sheet.title,
            "message": (
                f"Excel converti en CSV ({row_count} lignes, "
                f"feuille « {sheet.title} ») → {output}"
            ),
        }
    except (OSError, TypeError, ValueError, csv.Error) as exc:
        return {"error": str(exc)}
    finally:
        if workbook is not None:
            workbook.close()


async def read_xlsx(args: dict) -> dict:
    """Lit de façon bornée une feuille d'un classeur Excel."""
    workbook = None
    try:
        raw_path = str(args.get("path") or "").strip()
        if not raw_path:
            return {"error": "Paramètre 'path' requis."}
        requested_sheet = str(args.get("sheet") or "")
        action = str(args.get("action") or "read").lower()
        if action not in {"read", "sheets", "summary"}:
            return {"error": "Action invalide. Utilisez read, sheets ou summary."}
        max_rows = int(args.get("max_rows", 50))
        if not 1 <= max_rows <= _MAX_READ_ROWS:
            return {"error": f"'max_rows' doit être compris entre 1 et {_MAX_READ_ROWS}."}

        source = Path(_resolve_path(raw_path))
        if not source.is_file():
            return {"error": f"Fichier introuvable : {source}"}
        workbook = load_workbook(str(source), read_only=True, data_only=True)
        sheet_names = list(workbook.sheetnames)
        if action == "sheets":
            return {
                "success": True,
                "file": str(source),
                "sheets": sheet_names,
                "count": len(sheet_names),
            }

        sheet = (
            workbook[requested_sheet]
            if requested_sheet and requested_sheet in sheet_names
            else workbook.active
        )
        if action == "summary":
            return {
                "success": True,
                "file": str(source),
                "active_sheet": sheet.title,
                "sheets": sheet_names,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
            }

        headers: list[str] = []
        data: list[dict[str, str]] = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index == 0:
                headers = [
                    str(value) if value is not None else f"Col{column + 1}"
                    for column, value in enumerate(row)
                ]
            elif index <= max_rows:
                data.append(
                    {
                        headers[column]: str(value) if value is not None else ""
                        for column, value in enumerate(row)
                        if column < len(headers)
                    }
                )
            else:
                break
        return {
            "success": True,
            "file": str(source),
            "sheet": sheet.title,
            "headers": headers,
            "rows_loaded": len(data),
            "data": data,
        }
    except (OSError, TypeError, ValueError) as exc:
        return {"error": str(exc)}
    finally:
        if workbook is not None:
            workbook.close()
