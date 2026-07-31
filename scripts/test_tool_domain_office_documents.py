"""Régressions PowerPoint, extraction DOCX et validation Open XML."""

from __future__ import annotations

import asyncio
import sys
import tempfile
import zipfile
from pathlib import Path
from unittest.mock import patch


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

import tools
from agent_safety import validator_for
from office_packages import validate_office_package
from tool_domains import documents
from tool_domains import presentations


async def run_domain() -> None:
    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary)
        resolver = lambda value: str(Path(value).expanduser().resolve())
        presentations.configure(resolve_path=resolver)
        documents.configure(resolve_path=resolver)

        deck = root / "deck.pptx"
        created = await presentations.create(
            {
                "title": "Marvin",
                "subtitle": "Validation",
                "slides": [
                    {"title": "Architecture", "bullets": ["Runner", "Outils"]},
                    {"title": "Qualité", "points": ["Tests", "Rollback"]},
                ],
                "output_path": str(deck),
            }
        )
        assert created["success"] is True
        assert created["verified"] is True
        assert created["slides"] == 3
        assert created["style"] == "professional_dark"
        assert created["size_bytes"] == deck.stat().st_size
        assert len(created["sha256"]) == 64
        assert created["artifacts"] == [
            {
                "path": str(deck),
                "type": "presentation",
                "format": "pptx",
                "slides": 3,
            }
        ]
        validate_office_package(
            deck, suffix=".pptx", marker="ppt/presentation.xml"
        )
        read = await presentations.read({"path": str(deck)})
        assert read["slides"] == 3
        assert "Marvin" in read["text"]
        assert "Architecture" in read["text"]
        assert "Rollback" in read["text"]
        assert read["truncated"] is False
        directory_target = root / "exports"
        directory_target.mkdir()
        directory_created = await presentations.create(
            {"title": "Dossier", "output_path": str(directory_target)}
        )
        assert directory_created["success"] is True
        assert Path(directory_created["path"]).name == "Presentation_Marvin.pptx"
        assert "Style invalide" in (
            await presentations.create(
                {"title": "x", "style": "invented", "output_path": str(root / "x.pptx")}
            )
        )["error"]
        locked_target = root / "locked.pptx"
        original_create = presentations._create_presentation
        attempts = 0

        def fail_once(**kwargs):
            nonlocal attempts
            attempts += 1
            if attempts == 1:
                raise PermissionError("simulated PowerPoint lock")
            return original_create(**kwargs)

        with patch.object(presentations, "_create_presentation", side_effect=fail_once):
            locked_created = await presentations.create(
                {"title": "Verrou", "output_path": str(locked_target)}
            )
        assert locked_created["success"] is True
        assert locked_created["path"] != str(locked_target)
        assert "nouvelle-version" in Path(locked_created["path"]).stem
        assert Path(locked_created["path"]).is_file()

        from PIL import Image

        visual = root / "visual.jpg"
        Image.new("RGB", (1200, 800), (70, 96, 84)).save(visual)
        visual_deck = root / "visual-deck.pptx"
        visual_created = await presentations.create(
            {
                "title": "Aix",
                "slides": [{
                    "title": "Patrimoine",
                    "bullets": ["Une ville de culture", "Un territoire vivant"],
                    "image_path": str(visual),
                    "image_caption": "Source : site municipal",
                    "source_url": "https://example.test/visual.jpg",
                }],
                "output_path": str(visual_deck),
            }
        )
        assert visual_created["success"] is True
        with zipfile.ZipFile(visual_deck) as package:
            assert any(name.startswith("ppt/media/") for name in package.namelist())
        assert "extension .pptx" in (
            await presentations.create(
                {"title": "x", "output_path": str(root / "bad.txt")}
            )
        )["error"]
        assert "100 diapositives" in (
            await presentations.create(
                {
                    "slides": [{"title": str(index)} for index in range(101)],
                    "output_path": str(root / "too-many.pptx"),
                }
            )
        )["error"]

        from docx import Document

        source = root / "source.docx"
        document = Document()
        document.add_heading("Titre", level=1)
        document.add_paragraph("Avant")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Nom"
        table.cell(0, 1).text = "Valeur"
        table.cell(1, 0).text = "a|b"
        table.cell(1, 1).text = "42"
        document.add_paragraph("Après")
        document.save(source)

        output = root / "source.md"
        extracted = await documents.extract_docx_text(
            {
                "path": str(source),
                "format": "markdown",
                "output_path": str(output),
            }
        )
        assert extracted["success"] is True
        assert extracted["paragraphs"] == 3
        assert extracted["tables"] == 1
        full = output.read_text(encoding="utf-8")
        assert full.index("Avant") < full.index("| Nom | Valeur |") < full.index("Après")
        assert "a\\|b" in full
        assert full.startswith("# Titre")
        plain = await documents.extract_docx_text(
            {"path": str(source), "format": "text"}
        )
        assert "# Titre" not in plain["text"]
        assert "Nom | Valeur" in plain["text"]
        assert "sortie texte" in (
            await documents.extract_docx_text(
                {
                    "path": str(source),
                    "output_path": str(root / "bad.pdf"),
                }
            )
        )["error"]

        corrupt = root / "corrupt.pptx"
        corrupt.write_text("not a zip", encoding="utf-8")
        assert "Open XML" in (await presentations.read({"path": str(corrupt)}))["error"]

        bomb = root / "bomb.docx"
        with zipfile.ZipFile(bomb, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b"0" * 2_000_000)
        try:
            validate_office_package(
                bomb, suffix=".docx", marker="word/document.xml"
            )
        except ValueError as exc:
            assert "compression" in str(exc)
        else:
            raise AssertionError("Le taux de compression anormal doit être refusé.")


def main() -> int:
    assert tools.execute_pptx_create is presentations.create
    assert tools.execute_pptx_read is presentations.read
    assert tools.execute_docx_to_text is documents.extract_docx_text
    assert tools.TOOLS_REGISTRY["pptx_create"] is presentations.create
    assert tools.TOOLS_REGISTRY["pptx_read"] is presentations.read
    assert tools.TOOLS_REGISTRY["docx_to_text"] is documents.extract_docx_text

    validator = validator_for(tools.TOOLS_SPEC)
    assert validator(
        "pptx_create",
        {"slides": [{"title": "x", "unknown": True}]},
    )
    assert validator(
        "pptx_create",
        {"slides": [{"title": "x", "bullets": [42]}]},
    )
    asyncio.run(run_domain())
    print("Office Documents OK - PPTX, DOCX, ordre, atomicité et limites ZIP.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
