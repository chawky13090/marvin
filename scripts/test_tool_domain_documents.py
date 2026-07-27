"""Régressions du domaine DOCX et conversions, avec façade tools.py."""

from __future__ import annotations

import asyncio
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

import tools
from tool_domains import documents


async def main() -> int:
    with tempfile.TemporaryDirectory() as root:
        base = Path(root)
        documents.configure(resolve_path=lambda raw: str(base / raw))
        with patch.object(documents, "_document_identity", return_value={}):
            result = await documents.create_docx({
                "path": "report.docx",
                "title": "Rapport",
                "author": "Marvin",
                "content": (
                    "## Résultat\n\nTest **réussi**.\n\n"
                    "| Contrôle | État |\n|---|---|\n| API | OK |"
                ),
            })
        assert result["success"] is True
        output = base / "report.docx"
        assert output.exists() and output.stat().st_size > 0

        from docx import Document

        generated = Document(str(output))
        text = "\n".join(paragraph.text for paragraph in generated.paragraphs)
        assert "Rapport" in text and "Test réussi." in text
        assert generated.tables and generated.tables[0].cell(1, 1).text == "OK"

        wrong_source = base / "notes.txt"
        wrong_source.write_text("notes", encoding="utf-8")
        rejected = await documents.convert_docx_to_pdf({"path": "notes.txt"})
        assert ".docx" in rejected["error"]
        unsupported = await documents.convert_format({
            "path": "notes.txt", "target_format": "pdf"
        })
        assert "non supportée" in unsupported["error"]
        invalid_target = await documents.convert_format({
            "path": "report.docx", "target_format": "txt"
        })
        assert "pdf" in invalid_target["error"] and "docx" in invalid_target["error"]

        async def generate_text(_prompt, **_kwargs):
            return "# Document\n\n## Introduction\n\nContenu vérifié."

        documents.configure_generation(
            text_generator=generate_text,
            generated_dir=base / "generated",
        )
        with (
            patch.object(documents.db, "get_setting", return_value=None),
            patch.object(documents, "_document_identity", return_value={}),
        ):
            generated_result = await documents.generate({
                "title": "Procédure test",
                "text": "Notes",
                "template": "procedure",
                "author": "Marvin",
            })
        assert generated_result["success"] is True
        assert Path(generated_result["file"]).parent == base / "generated"
        assert Path(generated_result["file"]).exists()

    documents.configure(resolve_path=tools._resolve_path)
    assert tools.execute_file_create_docx is documents.create_docx
    assert tools.execute_file_convert_docx_to_pdf is documents.convert_docx_to_pdf
    assert tools.execute_document_convert_format is documents.convert_format
    assert tools.execute_document_generate is documents.generate
    print("Tool Domain Documents OK - DOCX réel, conversions et façade.")
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
